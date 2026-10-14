"""
siapkan_uji_nlu.py

Menyiapkan set uji NLU dari hasil kuesioner (ekspor Google Form) untuk models/evaluasi_nlu.py.

Tahap 1 -- buat lembar pelabelan:
    python models/siapkan_uji_nlu.py buat [data/kuisioner_NLU.csv]
  Setiap jawaban responden menjadi satu baris (responden x pertanyaan), tanpa timestamp. File
  kuesioner asli tidak diubah. Hasil: data/uji_nlu_label.xlsx dengan kolom tipe_kulit (dropdown),
  satu kolom per masalah kulit (isi 1 bila disebut), kolom "pakai" (jawaban yang tidak bisa berdiri
  sendiri, mis. "Gak", ditandai "tidak"), dan catatan. Lembar "Panduan" berisi aturan pelabelan.

Tahap 2 -- ambil hasil pelabelan:
    python models/siapkan_uji_nlu.py ambil
  Membaca data/uji_nlu_label.xlsx, membuang baris pakai = "tidak", memeriksa label, lalu menulis
  data/uji_nlu.csv (kolom id, kalimat, tipe_kulit, masalah_kulit, responden, pertanyaan, catatan).

Reliabilitas label (pelabel kedua, Cohen's kappa):
    python models/siapkan_uji_nlu.py anotator2
  Membuat data/uji_nlu_label_anotator2.xlsx: 50 kalimat acak (10 per pertanyaan, seed tetap) dengan
  format & panduan yang sama tetapi TANPA label pelabel pertama, untuk dilabeli orang kedua secara
  mandiri.
    python models/siapkan_uji_nlu.py kappa
  Membandingkan label kedua pelabel pada 50 kalimat itu: Cohen's kappa untuk tipe kulit (satu nilai,
  "kosong" dihitung sebagai kategori) dan untuk tiap masalah kulit (ada/tidak), persentase kesepakatan,
  serta kesepakatan persis per kalimat; ditafsirkan dengan skala Landis & Koch (1977). Ringkasan ditulis
  ke data/kappa_nlu.csv, kalimat yang labelnya berbeda ke data/kappa_nlu_selisih.csv untuk dibahas.

Pelabelan WAJIB dilakukan manusia: label buatan model (aturan atau LLM) akan menguntungkan
pendekatan yang memakai model tersebut, sehingga perbandingan tidak sah.
"""
import random
import re
import sys
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation
from sklearn.metrics import cohen_kappa_score

KUESIONER = Path("data/kuisioner_NLU.csv")
LEMBAR = Path("data/uji_nlu_label.xlsx")
HASIL = Path("data/uji_nlu.csv")
LEMBAR_ANOTATOR2 = Path("data/uji_nlu_label_anotator2.xlsx")
RINGKASAN_KAPPA = Path("data/kappa_nlu.csv")
SELISIH_KAPPA = Path("data/kappa_nlu_selisih.csv")
N_ANOTATOR2 = 50
SEED_ANOTATOR2 = 2026
TIPE = ["berminyak", "kering", "kombinasi", "sensitif", "normal"]
MASALAH = ["jerawat", "kusam", "hiperpigmentasi", "penuaan", "kemerahan_iritasi", "dehidrasi"]
# jawaban yang hanya ya/tidak -- tanpa pertanyaannya tidak bermakna sebagai pesan chat
_JAWAB_SINGKAT = {"gak", "ga", "nggak", "ngga", "enggak", "tidak", "tdk", "belum", "pernah", "iya", "ya", "-", "ada",
                  "tidak ada", "ga ada", "gak ada", "nggak ada", "belum pernah", "tidak pernah", "gak pernah"}

PANDUAN = [
    ("Kolom", "Isi"),
    ("tipe_kulit", "SATU dari: berminyak, kering, kombinasi, sensitif, normal. Kosongkan bila tidak disebut atau ragu."),
    ("kolom masalah", "Isi 1 pada setiap masalah yang disebut (boleh lebih dari satu); kosongkan yang tidak disebut."),
    ("pakai", "'tidak' untuk jawaban yang tidak bisa berdiri sendiri sebagai pesan chat (mis. 'Gak', 'Pernah')."),
    ("catatan", "Alasan bila ragu, atau kondisi di luar 6 label (pori besar, bopeng, mata panda, eksim)."),
    ("", ""),
    ("Kalau jawabannya menyebut...", "Label"),
    ("mengkilap, lengket, minyakan (seluruh wajah)", "tipe berminyak"),
    ("ketarik, kering, kasar, mengelupas, bersisik", "tipe kering"),
    ("berminyak di dahi/hidung (T-zone) tapi pipi normal/kering", "tipe kombinasi"),
    ("biasa aja, nyaman, tidak berminyak dan tidak kering", "tipe normal"),
    ("gampang/sering bereaksi, cocok-cocokan, mudah iritasi (kecenderungan)", "tipe sensitif"),
    ("merah, perih, gatal, ruam, panas, cekit-cekit (gejala yang dialami)", "masalah kemerahan_iritasi"),
    ("jerawat, bruntusan, komedo, breakout", "masalah jerawat"),
    ("bekas jerawat hitam/gelap/kecokelatan, flek, belang, warna tidak merata", "masalah hiperpigmentasi"),
    ("kusam, gelap, tidak cerah, kucel", "masalah kusam"),
    ("kerutan, garis halus, kendur", "masalah penuaan"),
    ("kurang lembap padahal berminyak, dehidrasi, kurang hidrasi", "masalah dehidrasi"),
    ("pori besar, bopeng, mata panda, eksim, skin barrier", "kosong (di luar 6 label) -- tulis di catatan"),
    ("", ""),
    ("Aturan", ""),
    ("Negasi", "Yang disangkal TIDAK dilabeli ('nggak berminyak' = bukan berminyak)."),
    ("Orang lain", "Kulit orang lain yang akan dibelikan tetap dilabeli."),
    ("Tanpa kondisi kulit", "Tetap dipakai dengan label kosong (menguji apakah pendekatan 'mengarang' label)."),
    ("Teks saja", "Label hanya dari teks jawaban, BUKAN dari pertanyaannya (sistem tidak melihat pertanyaan)."),
    ("Salah ketik", "Jangan diperbaiki; biarkan apa adanya."),
]
# Landis, J. R., & Koch, G. G. (1977). The measurement of observer agreement for categorical data. Biometrics.
_SKALA_KAPPA = [(0.20, "sangat lemah (slight)"), (0.40, "lemah (fair)"), (0.60, "sedang (moderate)"),
                (0.80, "kuat (substantial)"), (1.00, "hampir sempurna (almost perfect)")]


def _bersih(teks):
    return re.sub(r"\s+", " ", str(teks)).strip()


def _simpan_lembar(tabel, path, pertanyaan):
    """Tulis lembar pelabelan (sheet Label + Panduan + Pertanyaan) dengan dropdown & lebar kolom yang rapi."""
    with pd.ExcelWriter(path, engine="openpyxl") as w:
        tabel.to_excel(w, sheet_name="Label", index=False)
        pd.DataFrame(PANDUAN[1:], columns=PANDUAN[0]).to_excel(w, sheet_name="Panduan", index=False)
        pd.DataFrame({"kode": [f"P{i}" for i in range(1, len(pertanyaan) + 1)],
                      "pertanyaan": pertanyaan}).to_excel(w, sheet_name="Pertanyaan", index=False)

    wb = load_workbook(path)
    ws = wb["Label"]
    kepala = PatternFill("solid", fgColor="E8E6F8")
    for sel in ws[1]:
        sel.font, sel.fill = Font(bold=True), kepala
    lebar = {"A": 11, "B": 10, "C": 10, "D": 70, "E": 14, "L": 8, "M": 34}
    for kol, w_ in lebar.items():
        ws.column_dimensions[kol].width = w_
    for kol in "FGHIJK":
        ws.column_dimensions[kol].width = 11
    for row in ws.iter_rows(min_row=2):
        row[3].alignment = Alignment(wrap_text=True, vertical="top")
    n = ws.max_row
    dv_tipe = DataValidation(type="list", formula1='"' + ",".join(TIPE) + '"', allow_blank=True)
    dv_satu = DataValidation(type="list", formula1='"1"', allow_blank=True)
    dv_pakai = DataValidation(type="list", formula1='"ya,tidak"', allow_blank=False)
    for dv, rentang in ((dv_tipe, f"E2:E{n}"), (dv_satu, f"F2:K{n}"), (dv_pakai, f"L2:L{n}")):
        ws.add_data_validation(dv)
        dv.add(rentang)
    ws.freeze_panes = "E2"
    for nama in ("Panduan", "Pertanyaan"):
        lain = wb[nama]
        lain.column_dimensions["A"].width = 62
        lain.column_dimensions["B"].width = 100
        for sel in lain[1]:
            sel.font, sel.fill = Font(bold=True), kepala
    wb.save(path)


def _baris_kosong(id_, responden, pertanyaan, kalimat, pakai="ya", catatan=""):
    return {"id": id_, "responden": responden, "pertanyaan": pertanyaan, "kalimat": kalimat,
            "tipe_kulit": "", **{m: "" for m in MASALAH}, "pakai": pakai, "catatan": catatan}


def buat(path):
    df = pd.read_csv(path)
    pertanyaan = [c for c in df.columns if c.strip().lower() != "timestamp"]
    baris = []
    for r, (_, row) in enumerate(df.iterrows(), 1):
        for p, kolom in enumerate(pertanyaan, 1):
            teks = _bersih(row[kolom]) if pd.notna(row[kolom]) else ""
            if not teks:
                continue
            singkat = re.sub(r"[^a-z\- ]", "", teks.lower()).strip() in _JAWAB_SINGKAT
            baris.append(_baris_kosong(f"R{r:03d}-P{p}", f"R{r:03d}", f"P{p}", teks, "tidak" if singkat else "ya",
                                       "jawaban tidak bisa berdiri sendiri" if singkat else ""))
    tabel = pd.DataFrame(baris)
    _simpan_lembar(tabel, LEMBAR, [_bersih(c) for c in pertanyaan])
    print(f"{LEMBAR}: {len(tabel)} kalimat dari {len(df)} responden x {len(pertanyaan)} pertanyaan; "
          f"{int((tabel['pakai'] == 'tidak').sum())} ditandai 'tidak' (jawaban ya/tidak saja).")


def _baca_label(path):
    """Lembar pelabelan -> (DataFrame baris yang dipakai dengan kolom tipe & masalah (list), daftar kesalahan)."""
    df = pd.read_excel(path, sheet_name="Label", dtype=str).fillna("")
    df = df[df["pakai"].str.strip().str.lower() != "tidak"]
    salah, keluar = [], []
    for _, r in df.iterrows():
        tipe = r["tipe_kulit"].strip().lower()
        if tipe and tipe not in TIPE:
            salah.append(f"{r['id']}: tipe_kulit '{tipe}' tidak dikenal")
        aneh = [m for m in MASALAH if r[m].strip() not in ("", "0", "1")]
        if aneh:
            salah.append(f"{r['id']}: kolom {aneh} harus 1 atau kosong")
        keluar.append({"id": r["id"], "kalimat": r["kalimat"], "tipe_kulit": tipe,
                       "masalah": [m for m in MASALAH if r[m].strip() not in ("", "0")],
                       "responden": r["responden"], "pertanyaan": r["pertanyaan"], "catatan": r["catatan"]})
    return pd.DataFrame(keluar), [f"{path.name} {s}" for s in salah]


def ambil():
    label, salah = _baca_label(LEMBAR)
    if salah:
        sys.exit("Label belum valid:\n  " + "\n  ".join(salah))
    hasil = label.assign(masalah_kulit=label["masalah"].map(";".join)).drop(columns="masalah")[
        ["id", "kalimat", "tipe_kulit", "masalah_kulit", "responden", "pertanyaan", "catatan"]]
    hasil.to_csv(HASIL, index=False)
    berlabel = ((hasil["tipe_kulit"] != "") | (hasil["masalah_kulit"] != "")).sum()
    print(f"{HASIL}: {len(hasil)} kalimat ({berlabel} memiliki label tipe/masalah, {len(hasil) - berlabel} tanpa label).")


def anotator2():
    """50 kalimat acak (sama banyak per pertanyaan) untuk pelabel kedua, tanpa label pelabel pertama."""
    if LEMBAR_ANOTATOR2.exists():
        sys.exit(f"{LEMBAR_ANOTATOR2} sudah ada (mungkin sudah mulai dilabeli). Hapus dulu kalau memang mau dibuat ulang.")
    df = pd.read_excel(LEMBAR, sheet_name="Label", dtype=str).fillna("")
    df = df[df["pakai"].str.strip().str.lower() != "tidak"]
    acak = random.Random(SEED_ANOTATOR2)
    kelompok = {p: sorted(g["id"]) for p, g in df.groupby("pertanyaan")}
    per = N_ANOTATOR2 // len(kelompok)
    pilih = [i for ids in kelompok.values() for i in acak.sample(ids, min(per, len(ids)))]
    sisa = sorted(set(df["id"]) - set(pilih))
    pilih += acak.sample(sisa, max(0, N_ANOTATOR2 - len(pilih)))
    sampel = df[df["id"].isin(pilih)].sort_values("id")
    # hanya teks yang disalin -- label pelabel pertama tidak boleh terlihat oleh pelabel kedua
    tabel = pd.DataFrame([_baris_kosong(r["id"], r["responden"], r["pertanyaan"], r["kalimat"])
                          for _, r in sampel.iterrows()])
    pertanyaan = pd.read_excel(LEMBAR, sheet_name="Pertanyaan", dtype=str)["pertanyaan"].tolist()
    _simpan_lembar(tabel, LEMBAR_ANOTATOR2, pertanyaan)
    print(f"{LEMBAR_ANOTATOR2}: {len(tabel)} kalimat "
          f"({', '.join(f'{p}={n}' for p, n in tabel['pertanyaan'].value_counts().sort_index().items())}), seed "
          f"{SEED_ANOTATOR2}. Berikan file ini ke pelabel kedua bersama sheet Panduan; jangan perlihatkan labelmu.")


def tafsir_kappa(k):
    if k != k:
        return "tidak terdefinisi"
    if k < 0:
        return "buruk (poor)"
    return next(teks for batas, teks in _SKALA_KAPPA if k <= batas)


def _kappa(a, b):
    """Cohen's kappa; tidak terdefinisi (NaN) kalau kedua pelabel hanya memakai satu nilai yang sama."""
    if len(set(a) | set(b)) < 2:
        return float("nan")
    return float(cohen_kappa_score(a, b))


def kappa():
    if not LEMBAR_ANOTATOR2.exists():
        sys.exit(f"{LEMBAR_ANOTATOR2} belum ada. Buat dulu: python models/siapkan_uji_nlu.py anotator2")
    satu, salah1 = _baca_label(LEMBAR)
    dua, salah2 = _baca_label(LEMBAR_ANOTATOR2)
    if salah1 or salah2:
        sys.exit("Label belum valid:\n  " + "\n  ".join(salah1 + salah2))
    g = satu.merge(dua, on="id", suffixes=("_1", "_2"))
    if g.empty:
        sys.exit("Tidak ada kalimat yang sama di kedua lembar.")
    n = len(g)

    baris = []
    t1 = [t or "(kosong)" for t in g["tipe_kulit_1"]]
    t2 = [t or "(kosong)" for t in g["tipe_kulit_2"]]
    k = _kappa(t1, t2)
    baris.append({"label": "tipe kulit (6 kategori termasuk kosong)", "n": n,
                  "dipakai_pelabel_1": sum(t != "(kosong)" for t in t1), "dipakai_pelabel_2": sum(t != "(kosong)" for t in t2),
                  "kesepakatan_persen": round(100 * sum(x == y for x, y in zip(t1, t2)) / n, 1),
                  "kappa": round(k, 3), "tafsiran": tafsir_kappa(k)})
    kappa_masalah = []
    for m in MASALAH:
        a = [int(m in x) for x in g["masalah_1"]]
        b = [int(m in x) for x in g["masalah_2"]]
        k = _kappa(a, b)
        if k == k:
            kappa_masalah.append(k)
        baris.append({"label": f"masalah: {m}", "n": n, "dipakai_pelabel_1": sum(a), "dipakai_pelabel_2": sum(b),
                      "kesepakatan_persen": round(100 * sum(x == y for x, y in zip(a, b)) / n, 1),
                      "kappa": round(k, 3), "tafsiran": tafsir_kappa(k)})
    rata = sum(kappa_masalah) / len(kappa_masalah) if kappa_masalah else float("nan")
    baris.append({"label": "masalah kulit (rata-rata kappa label yang terdefinisi)", "n": n,
                  "dipakai_pelabel_1": "", "dipakai_pelabel_2": "", "kesepakatan_persen": "",
                  "kappa": round(rata, 3), "tafsiran": tafsir_kappa(rata)})
    persis = [(x["tipe_kulit_1"] == x["tipe_kulit_2"]) and (set(x["masalah_1"]) == set(x["masalah_2"]))
              for _, x in g.iterrows()]
    baris.append({"label": "kesepakatan persis per kalimat (tipe & semua masalah)", "n": n, "dipakai_pelabel_1": "",
                  "dipakai_pelabel_2": "", "kesepakatan_persen": round(100 * sum(persis) / n, 1), "kappa": "",
                  "tafsiran": ""})
    ringkas = pd.DataFrame(baris)
    ringkas.to_csv(RINGKASAN_KAPPA, index=False)

    beda = g[[not p for p in persis]]
    pd.DataFrame({"id": beda["id"], "kalimat": beda["kalimat_1"],
                  "tipe_pelabel_1": beda["tipe_kulit_1"], "tipe_pelabel_2": beda["tipe_kulit_2"],
                  "masalah_pelabel_1": beda["masalah_1"].map(";".join),
                  "masalah_pelabel_2": beda["masalah_2"].map(";".join)}).to_csv(SELISIH_KAPPA, index=False)

    with pd.option_context("display.max_columns", None, "display.width", 200, "display.max_colwidth", 60):
        print(ringkas.to_string(index=False))
    print(f"\n{RINGKASAN_KAPPA}: ringkasan untuk Bab IV; {SELISIH_KAPPA}: {len(beda)} kalimat yang labelnya berbeda "
          "(bahas berdua, lalu perbaiki lembar utama bila kalian sepakat labelmu yang keliru).")


if __name__ == "__main__":
    perintah = {"buat": lambda: buat(Path(sys.argv[2]) if len(sys.argv) > 2 else KUESIONER), "ambil": ambil,
                "anotator2": anotator2, "kappa": kappa}
    if len(sys.argv) < 2 or sys.argv[1] not in perintah:
        sys.exit("Pakai: python models/siapkan_uji_nlu.py buat [kuesioner.csv] | ambil | anotator2 | kappa")
    perintah[sys.argv[1]]()
